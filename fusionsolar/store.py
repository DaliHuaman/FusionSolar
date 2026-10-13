"""Persistencia: CSV maestro local (siempre) + Excel historico en OneDrive (con cola si esta bloqueado)."""
from __future__ import annotations

import csv
import logging
import os
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .common import COLUMNAS, STATE_DIR

log = logging.getLogger("fusionsolar.store")

CSV_MAESTRO = STATE_DIR / "historico.csv"
PENDIENTES = STATE_DIR / "pendientes_excel.csv"
LIMITE_FILAS = 1_000_000          # Excel admite 1.048.576; dejamos margen
HOJA = "Historico"
FMT = "%Y-%m-%d %H:%M:%S"


# --------------------------------------------------------------------------- CSV
def _leer_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def _escribir_csv(path: Path, filas: list[dict], anexar: bool) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    nuevo = not path.exists() or not anexar
    with open(path, "a" if anexar else "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNAS)
        if nuevo:
            w.writeheader()
        w.writerows(filas)


def anexar_csv(filas: list[dict]) -> None:
    _escribir_csv(CSV_MAESTRO, filas, anexar=True)


def construir_fila(ts: datetime, inv: dict, lectura: dict, estado: str, fuente: str = "FusionSolar") -> dict:
    r = lectura.get("resistencia")
    um = inv["umbral"]
    return {
        "Fecha_Hora": ts.strftime(FMT),
        "Fecha": ts.strftime("%Y-%m-%d"),
        "Hora": ts.strftime("%H:%M"),
        "Planta": inv["planta"],
        "Grupo": inv["grupo"],
        "Inversor": inv["nombre"],
        "Resistencia_MOhm": "" if r is None else r,
        "Umbral_MOhm": um,
        "Margen_%": "" if r is None else round((r - um) / um * 100, 1),
        "Estado_Aislamiento": estado,
        "Estado_Potencia_Activa": lectura.get("potencia_txt", ""),
        "Potencia_Activa_%": "" if lectura.get("potencia_pct") is None else lectura["potencia_pct"],
        "Estado_Inversor": lectura.get("estado_inversor", ""),
        "Fuente": fuente,
    }


# --------------------------------------------------------------------------- Excel
def _ruta_excel() -> Path:
    base = Path(os.environ.get("EXCEL_DIR") or (STATE_DIR / "excel"))
    base.mkdir(parents=True, exist_ok=True)
    nombre = os.environ.get("EXCEL_NOMBRE", "Historico_Indicadores_Inversores.xlsx")
    stem, ext = os.path.splitext(nombre)
    parte = 1
    while True:
        p = base / (nombre if parte == 1 else f"{stem}_parte{parte}{ext}")
        if not p.exists():
            return p
        wb = load_workbook(p, read_only=True)
        n = wb[HOJA].max_row if HOJA in wb.sheetnames else 0
        wb.close()
        if n < LIMITE_FILAS:
            return p
        parte += 1


def _valor_celda(col: str, v: str):
    if v == "" or v is None:
        return None
    if col == "Fecha_Hora":
        return datetime.strptime(v, FMT)
    if col == "Fecha":
        return datetime.strptime(v, "%Y-%m-%d")
    if col in ("Resistencia_MOhm", "Umbral_MOhm", "Margen_%", "Potencia_Activa_%"):
        return float(v)
    return v


def _nuevo_libro(umbrales: list[dict]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    ws.append(COLUMNAS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    anchos = [19, 11, 7, 8, 11, 10, 15, 13, 10, 18, 52, 14, 22, 12]
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    # formato condicional por estado (columna J)
    rng = "J2:J1048576"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ALARMA"'],
                                  fill=PatternFill("solid", bgColor="F8CBAD"), font=Font(color="9C0006", bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ADVERTENCIA"'],
                                  fill=PatternFill("solid", bgColor="FFE699"), font=Font(color="7F6000")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'],
                                  fill=PatternFill("solid", bgColor="C6EFCE"), font=Font(color="006100")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"SIN_DATO"'],
                                  fill=PatternFill("solid", bgColor="EDEDED"), font=Font(color="595959")))
    # hoja de umbrales
    wu = wb.create_sheet("Umbrales")
    wu.append(["Planta", "Grupo", "Inversor", "Umbral resistencia aislamiento (MΩ)"])
    for c in wu[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", bgColor="1F4E79", fgColor="1F4E79")
    for u in umbrales:
        wu.append([u["planta"], u["grupo"], u["nombre"], u["umbral"]])
    for i, w in enumerate([10, 12, 12, 38], 1):
        wu.column_dimensions[get_column_letter(i)].width = w
    wb.create_sheet("Resumen")
    return wb


def _resumen(wb: Workbook, ultimas: list[dict]) -> None:
    ws = wb["Resumen"]
    ws.delete_rows(1, ws.max_row)
    ws.append(["Última lectura por inversor (se reescribe en cada corrida)"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Fecha_Hora", "Planta", "Grupo", "Inversor", "Resistencia_MOhm", "Umbral_MOhm",
               "Estado_Aislamiento", "Estado_Potencia_Activa"])
    for c in ws[2]:
        c.font = Font(bold=True)
    for f in ultimas:
        ws.append([f["Fecha_Hora"], f["Planta"], f["Grupo"], f["Inversor"],
                   _valor_celda("Resistencia_MOhm", f["Resistencia_MOhm"]), float(f["Umbral_MOhm"]),
                   f["Estado_Aislamiento"], f["Estado_Potencia_Activa"]])
    for i, w in enumerate([19, 8, 11, 10, 16, 13, 18, 52], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def escribir_excel(filas: list[dict], umbrales: list[dict]) -> Path:
    """Anexa filas al Excel historico. Lanza PermissionError/OSError si el archivo esta bloqueado."""
    ruta = _ruta_excel()
    wb = load_workbook(ruta) if ruta.exists() else _nuevo_libro(umbrales)
    ws = wb[HOJA]
    for f in filas:
        ws.append([_valor_celda(c, f.get(c, "")) for c in COLUMNAS])
        r = ws.max_row
        ws.cell(r, 1).number_format = "yyyy-mm-dd hh:mm:ss"
        ws.cell(r, 2).number_format = "yyyy-mm-dd"
        ws.cell(r, 7).number_format = "0.000"
        ws.cell(r, 8).number_format = "0.000"
    ref = f"A1:{get_column_letter(len(COLUMNAS))}{ws.max_row}"
    if ws.tables:
        ws.tables["tHistorico"].ref = ref
        ws.tables["tHistorico"].autoFilter.ref = ref
    else:
        t = Table(displayName="tHistorico", ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(t)
    # Resumen: ultima lectura de cada inversor entre las filas que acabamos de escribir
    ult: dict[tuple, dict] = {}
    for f in filas:
        ult[(f["Planta"], f["Grupo"], f["Inversor"])] = f
    if "Resumen" in wb.sheetnames and ult:
        _resumen(wb, list(ult.values()))
    tmp = ruta.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, ruta)          # reemplazo atomico; falla con PermissionError si Excel lo tiene abierto
    return ruta


def guardar(filas: list[dict], umbrales: list[dict]) -> dict:
    """CSV maestro siempre; Excel con cola de pendientes si esta bloqueado."""
    anexar_csv(filas)
    pend = _leer_csv(PENDIENTES)
    total = pend + filas
    try:
        ruta = escribir_excel(total, umbrales)
        if PENDIENTES.exists():
            PENDIENTES.unlink()
        return {"excel": str(ruta), "pendientes": 0}
    except (PermissionError, OSError) as e:
        log.warning("Excel bloqueado/no disponible (%s). %s filas quedan pendientes.", e, len(total))
        _escribir_csv(PENDIENTES, total, anexar=False)
        return {"excel": None, "pendientes": len(total), "error": str(e)}


def leer_historico() -> list[dict]:
    return _leer_csv(CSV_MAESTRO)
