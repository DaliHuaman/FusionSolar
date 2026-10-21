"""Escritura en la plantilla "Indicadores_FusionSolar.xlsx".

Hoja "Indicadores" (la tuya): una fila por corrida, hora en A y pares MΩ / % por inversor, con una
columna vacía separando cada grupo:
    B:E LC (Inv1, Inv2) | F | G:J VES (Inv1, Inv2) | K | L:Q PL Planta A (A1-A3) | R | S:X PL Planta B (B1-B3)
    La tabla de umbrales queda en Z:AA.
Hoja "Historico": formato plano, una fila por lectura e inversor. Es la mejor fuente para Power BI.

Hoja "BESS" (segunda hoja): una fila por corrida y un bloque de 6 temperaturas (°C) por BESS, separados por una
columna vacía. Se reconstruye entera en cada corrida desde el CSV BESS (acotado a los últimos N días), así que
agregarla no cambia el diseño de las otras hojas ni dispara la reconstrucción del archivo.

Si el archivo existente tiene un diseño anterior (otra VERSION_DISENO), se reconstruye desde el CSV maestro.
"""
from __future__ import annotations

import os
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Callable

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .common import CAMPOS_BESS, COLUMNAS, ROOT

HOJA_WIDE = "Indicadores"
HOJA_PLANA = "Historico"
PLANTILLA = ROOT / "plantilla" / "Indicadores_FusionSolar.xlsx"
VERSION_DISENO = "layout-2"
HOJA_BESS = "BESS"
TITULOS_BESS = ["Máx. grupo", "Agua sum. batería", "Agua ret. batería", "Agua sum. potencia", "Agua ret. potencia", "Exterior"]
LIMITE_FILAS = 200_000      # ~14 meses de histórico plano (10 inversores x 48 lecturas/día); luego se crea _parte2
FMT = "%Y-%m-%d %H:%M:%S"
PRIMERA_FILA_DATOS = 4

# inversor -> columna de MΩ en la hoja ancha (el % va en la siguiente)
COLS = {"LC-Inv1": 2, "LC-Inv2": 4, "VES-Inv1": 7, "VES-Inv2": 9,
        "PL-InvA.1": 12, "PL-InvA.2": 14, "PL-InvA.3": 16, "PL-InvB.1": 19, "PL-InvB.2": 21, "PL-InvB.3": 23}
ORDEN = list(COLS)          # el mismo orden de la tabla de umbrales (Z5:AA14)
NOMBRES = {"LC-Inv1": "INVERSOR 1", "LC-Inv2": "INVERSOR 2", "VES-Inv1": "INVERSOR 1", "VES-Inv2": "INVERSOR 2",
           "PL-InvA.1": "INVERSOR A1", "PL-InvA.2": "INVERSOR A2", "PL-InvA.3": "INVERSOR A3",
           "PL-InvB.1": "INVERSOR B1", "PL-InvB.2": "INVERSOR B2", "PL-InvB.3": "INVERSOR B3"}
ETIQ_UMBRAL = ["Inv1 LC", "Inv2 LC", "Inv1 VES", "Inv2 VES", "Inv A1", "Inv A2", "Inv A3", "Inv B1", "Inv B2", "Inv B3"]
GRUPOS = [("BSF LOGISCITY", 2, 5), ("BSF VILLA EL SALVADOR", 7, 10), ("PLANTA SOLAR Y BESS - BSF PL", 12, 24)]
SEPARADORES = (6, 11, 18)   # F, K, R
COL_UMB = 26                # Z (etiqueta) y AA (valor)
FILA_UMB0 = 5


def _estilo(dst, src) -> None:
    dst._style = copy(src._style)


def _preparar(wb: Workbook, umbrales: dict[str, float]) -> None:
    """Adapta la plantilla de origen al diseño con separadores y agrega la hoja plana."""
    ws = wb[HOJA_WIDE]
    g1, g2, g3 = ws["B1"], ws["B2"], ws["B3"]
    estilo = {k: copy(c._style) for k, c in (("g1", g1), ("g2", g2), ("g3", g3))}
    neutro = copy(ws.cell(1, 40)._style)                       # celda sin formato
    # copia de la tabla de umbrales (T1:U14) antes de limpiar
    umb = [[(ws.cell(r, c).value, copy(ws.cell(r, c)._style)) for c in (20, 21)] for r in range(1, 15)]
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))
    for r in range(1, 15):
        for c in range(2, 30):
            cel = ws.cell(r, c)
            cel.value = None
            cel._style = copy(neutro)
    # encabezados por grupo
    for nombre, c0, c1 in GRUPOS:
        for c in range(c0, c1 + 1):
            ws.cell(1, c)._style = copy(estilo["g1"])
        ws.cell(1, c0).value = nombre
        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c1)
    for iid, c in COLS.items():
        for k in (0, 1):
            ws.cell(2, c + k)._style = copy(estilo["g2"])
            ws.cell(3, c + k)._style = copy(estilo["g3"])
        ws.cell(2, c).value = NOMBRES[iid]
        ws.merge_cells(start_row=2, start_column=c, end_row=2, end_column=c + 1)
        ws.cell(3, c).value, ws.cell(3, c + 1).value = "MΩ", "%"
        ws.column_dimensions[get_column_letter(c)].width = 9
        ws.column_dimensions[get_column_letter(c + 1)].width = 7
    for c in SEPARADORES:
        ws.column_dimensions[get_column_letter(c)].width = 2.5
    # tabla de umbrales en Z:AA (leyenda de filas 1-2 y título de la fila 4 incluidos)
    for r, par in enumerate(umb, start=1):
        for k, (valor, est) in enumerate(par):
            cel = ws.cell(r, COL_UMB + k)
            cel.value, cel._style = valor, copy(est)
    ws.merge_cells(start_row=4, start_column=COL_UMB, end_row=4, end_column=COL_UMB + 1)
    for k, iid in enumerate(ORDEN):
        ws.cell(FILA_UMB0 + k, COL_UMB).value = ETIQ_UMBRAL[k]
        ws.cell(FILA_UMB0 + k, COL_UMB + 1).value = umbrales.get(iid)
    ws.column_dimensions[get_column_letter(COL_UMB)].width = 30
    ws.column_dimensions[get_column_letter(COL_UMB + 1)].width = 28
    ws.column_dimensions["A"].width = 17
    ws.freeze_panes = "B4"
    # MΩ bajo el umbral en rojo (las celdas vacías no se pintan)
    rojo = PatternFill("solid", bgColor="F8CBAD")
    for k, iid in enumerate(ORDEN):
        col = get_column_letter(COLS[iid])
        ref = f"${get_column_letter(COL_UMB + 1)}${FILA_UMB0 + k}"
        ws.conditional_formatting.add(
            f"{col}{PRIMERA_FILA_DATOS}:{col}1048576",
            FormulaRule(formula=[f"AND(ISNUMBER({col}{PRIMERA_FILA_DATOS}),{col}{PRIMERA_FILA_DATOS}<{ref})"],
                        fill=rojo, font=Font(color="9C0006", bold=True)))
    # hoja plana
    wp = wb.create_sheet(HOJA_PLANA)
    wp.append(COLUMNAS)
    for c in wp[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wp.freeze_panes = "A2"
    for i, w in enumerate([19, 11, 7, 8, 11, 10, 15, 13, 10, 18, 52, 14, 22, 14], 1):
        wp.column_dimensions[get_column_letter(i)].width = w
    for txt, bg, fg in (("ALARMA", "F8CBAD", "9C0006"), ("ADVERTENCIA", "FFE699", "7F6000"),
                        ("OK", "C6EFCE", "006100"), ("SIN_DATO", "EDEDED", "595959")):
        wp.conditional_formatting.add("J2:J1048576", CellIsRule(
            operator="equal", formula=[f'"{txt}"'], fill=PatternFill("solid", bgColor=bg), font=Font(color=fg)))
    wb.properties.keywords = VERSION_DISENO


def _num(v):
    return None if v in ("", None) else float(v)


def _valor_plano(col: str, v):
    if v in ("", None):
        return None
    if col == "Fecha_Hora":
        return datetime.strptime(v, FMT)
    if col == "Fecha":
        return datetime.strptime(v, "%Y-%m-%d")
    if col in ("Resistencia_MOhm", "Umbral_MOhm", "Margen_%", "Potencia_Activa_%"):
        return float(v)
    return v


def ruta_excel() -> Path:
    base = Path(os.environ.get("EXCEL_DIR") or (ROOT / "state" / "excel"))
    base.mkdir(parents=True, exist_ok=True)
    nombre = os.environ.get("EXCEL_NOMBRE", "Indicadores_FusionSolar.xlsx")
    stem, ext = os.path.splitext(nombre)
    parte = 1
    while True:
        p = base / (nombre if parte == 1 else f"{stem}_parte{parte}{ext}")
        if not p.exists():
            return p
        wb = load_workbook(p, read_only=True)
        n = wb[HOJA_PLANA].max_row if HOJA_PLANA in wb.sheetnames else 0
        wb.close()
        if n < LIMITE_FILAS:
            return p
        parte += 1


def _agregar(wb: Workbook, filas: list[dict]) -> None:
    wp = wb[HOJA_PLANA]
    for f in filas:
        wp.append([_valor_plano(c, f.get(c)) for c in COLUMNAS])
        r = wp.max_row
        wp.cell(r, 1).number_format = "yyyy-mm-dd hh:mm:ss"
        wp.cell(r, 2).number_format = "yyyy-mm-dd"
        wp.cell(r, 7).number_format = wp.cell(r, 8).number_format = "0.000"
    ref = f"A1:{get_column_letter(len(COLUMNAS))}{wp.max_row}"
    if wp.tables:
        wp.tables["tHistorico"].ref = ref
        wp.tables["tHistorico"].autoFilter.ref = ref
    else:
        t = Table(displayName="tHistorico", ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        wp.add_table(t)

    ws = wb[HOJA_WIDE]
    por_ts: dict[str, dict[str, dict]] = {}
    for f in filas:
        por_ts.setdefault(f["Fecha_Hora"], {})[f"{f['Planta']}-{f['Inversor']}"] = f
    fila = PRIMERA_FILA_DATOS
    while ws.cell(fila, 1).value is not None:
        fila += 1
    for ts in sorted(por_ts):
        ws.cell(fila, 1).value = datetime.strptime(ts, FMT).replace(second=0)
        ws.cell(fila, 1).number_format = "dd/mm/yyyy hh:mm"
        for iid, f in por_ts[ts].items():
            c = COLS.get(iid)
            if not c:
                continue
            ws.cell(fila, c).value = _num(f["Resistencia_MOhm"])
            ws.cell(fila, c).number_format = "0.000"
            ws.cell(fila, c + 1).value = _num(f["Potencia_Activa_%"])
            ws.cell(fila, c + 1).number_format = "0.0"
        fila += 1


def _hoja_bess(wb: Workbook, filas_bess: list[dict]) -> None:
    """(Re)crea la hoja "BESS" como segunda hoja: A = horario, luego 6 columnas °C por BESS separadas por una vacía."""
    if HOJA_BESS in wb.sheetnames:
        del wb[HOJA_BESS]
    ws = wb.create_sheet(HOJA_BESS, 1)
    ids = sorted({f["BESS"] for f in filas_bess})
    cols = list(CAMPOS_BESS.values())
    azul = PatternFill("solid", fgColor="1F4E79")
    gris = PatternFill("solid", fgColor="DDEBF7")
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for fila in (1, 2):
        ws.cell(fila, 1).fill = azul if fila == 1 else gris
    ws.cell(1, 1).value, ws.cell(2, 1).value = "Horario", "°C"
    ws.cell(1, 1).font = Font(bold=True, color="FFFFFF")
    ws.cell(1, 1).alignment = ws.cell(2, 1).alignment = centro
    col0 = {}
    for k, bid in enumerate(ids):
        c = 2 + k * (len(cols) + 1)
        col0[bid] = c
        for j, tit in enumerate(TITULOS_BESS):
            ws.cell(1, c + j).fill = azul
            ws.cell(2, c + j).value, ws.cell(2, c + j).fill, ws.cell(2, c + j).alignment = tit, gris, centro
            ws.column_dimensions[get_column_letter(c + j)].width = 11
        ws.cell(1, c).value = bid
        ws.cell(1, c).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, c).alignment = centro
        ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + len(cols) - 1)
        if k:
            ws.column_dimensions[get_column_letter(c - 1)].width = 2.5
    ws.column_dimensions["A"].width = 17
    ws.row_dimensions[2].height = 30
    por_ts: dict[str, dict[str, dict]] = {}
    for f in filas_bess:
        por_ts.setdefault(f["Fecha_Hora"], {})[f["BESS"]] = f
    for fila, ts in enumerate(sorted(por_ts), start=3):
        a = ws.cell(fila, 1)
        a.value, a.number_format = datetime.strptime(ts, FMT).replace(second=0), "dd/mm/yyyy hh:mm"
        for bid, f in por_ts[ts].items():
            for j, col in enumerate(cols):
                cel = ws.cell(fila, col0[bid] + j)
                cel.value, cel.number_format = _num(f.get(col)), "0.0"
    ws.freeze_panes = "B3"


def escribir(filas: list[dict], inversores: list[dict],
             historico_completo: Callable[[], list[dict]] | None = None,
             filas_bess: list[dict] | None = None) -> Path:
    """Anexa las lecturas al Excel. Lanza PermissionError/OSError si el archivo está bloqueado.

    `historico_completo` devuelve todas las filas del CSV maestro (que ya incluye `filas`): se usa solo
    para reconstruir el Excel cuando el archivo existente tiene un diseño antiguo.
    """
    umbrales = {i["id"]: i["umbral"] for i in inversores}
    ruta = ruta_excel()
    wb, a_escribir = None, filas
    if ruta.exists():
        wb = load_workbook(ruta)
        if wb.properties.keywords != VERSION_DISENO:          # diseño antiguo -> reconstruir desde el CSV
            wb = None
            if historico_completo:
                a_escribir = sorted(historico_completo(), key=lambda f: f["Fecha_Hora"])
    if wb is None:
        if not PLANTILLA.exists():
            raise FileNotFoundError(f"No existe la plantilla {PLANTILLA}")
        wb = load_workbook(PLANTILLA)
        _preparar(wb, umbrales)
    _agregar(wb, a_escribir)
    if filas_bess:
        _hoja_bess(wb, filas_bess)
    tmp = ruta.with_name(ruta.stem + ".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, ruta)        # atómico; falla con PermissionError si Excel tiene el archivo abierto
    return ruta
